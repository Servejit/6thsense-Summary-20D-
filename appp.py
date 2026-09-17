```python
import streamlit as st
from openpyxl import load_workbook, Workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter
from copy import copy
import re
import io


# ============================================================
# PAGE
# ============================================================

st.set_page_config(
    page_title="6thsense Vardaan Auto-Color",
    page_icon="📊",
    layout="wide"
)

st.title("6thsense Vardaan Auto-Color")
st.write(
    "Upload the Excel file and generate the exact auto-coloured output."
)


# ============================================================
# UPLOAD
# ============================================================

uploaded_file = st.file_uploader(
    "Upload Excel file",
    type=["xlsx", "xlsm"]
)


if uploaded_file is not None:

    if st.button(
        "Generate Auto-Color Excel",
        type="primary"
    ):

        try:

            # ==================================================
            # READ UPLOADED FILE
            # ==================================================

            file_bytes = uploaded_file.getvalue()

            wb_values = load_workbook(
                io.BytesIO(file_bytes),
                data_only=True
            )

            wb_format = load_workbook(
                io.BytesIO(file_bytes),
                data_only=False
            )

            if "summary" not in wb_values.sheetnames:

                st.error(
                    "Sheet 'summary' not found in the uploaded file."
                )

                st.stop()

            src = wb_values["summary"]
            src_format = wb_format["summary"]


            # ==================================================
            # CREATE NEW WORKBOOK
            # ==================================================

            wb = Workbook()

            ws = wb.active
            ws.title = "summary"


            # ==================================================
            # COPY VALUES ONLY + ORIGINAL FORMATTING
            # ==================================================

            for r in range(
                1,
                src.max_row + 1
            ):

                for c in range(
                    1,
                    src.max_column + 1
                ):

                    # ------------------------------------------
                    # VALUE ONLY - NO FORMULA
                    # ------------------------------------------

                    ws.cell(
                        r,
                        c
                    ).value = src.cell(
                        r,
                        c
                    ).value


                    # ------------------------------------------
                    # ORIGINAL FORMATTING
                    # ------------------------------------------

                    old = src_format.cell(
                        r,
                        c
                    )

                    new = ws.cell(
                        r,
                        c
                    )

                    if old.has_style:

                        new._style = copy(
                            old._style
                        )

                    new.number_format = (
                        old.number_format
                    )

                    new.font = copy(
                        old.font
                    )

                    new.fill = copy(
                        old.fill
                    )

                    new.border = copy(
                        old.border
                    )

                    new.alignment = copy(
                        old.alignment
                    )

                    new.protection = copy(
                        old.protection
                    )


            # ==================================================
            # COPY COLUMN WIDTHS
            # ==================================================

            for col, dim in (
                src_format.column_dimensions.items()
            ):

                ws.column_dimensions[
                    col
                ].width = dim.width

                ws.column_dimensions[
                    col
                ].hidden = dim.hidden


            # ==================================================
            # COPY ROW HEIGHTS
            # ==================================================

            for row, dim in (
                src_format.row_dimensions.items()
            ):

                ws.row_dimensions[
                    row
                ].height = dim.height

                ws.row_dimensions[
                    row
                ].hidden = dim.hidden


            # ==================================================
            # COPY MERGED CELLS
            # ==================================================

            for merged in (
                src_format.merged_cells.ranges
            ):

                ws.merge_cells(
                    str(merged)
                )


            # ==================================================
            # FIND HEADER ROW
            # ==================================================

            header_row = None

            for r in range(
                1,
                ws.max_row + 1
            ):

                value = ws.cell(
                    r,
                    1
                ).value

                if value is not None:

                    if (
                        str(value)
                        .strip()
                        .lower()
                        == "symbol"
                    ):

                        header_row = r
                        break


            if header_row is None:

                st.error(
                    "Header row containing 'Symbol' was not found."
                )

                st.stop()


            # ==================================================
            # CLEAN TEXT
            # ==================================================

            def clean_text(value):

                if value is None:
                    return ""

                text = str(value)

                text = text.replace(
                    "\xa0",
                    " "
                )

                text = text.replace(
                    "\u200b",
                    ""
                )

                text = re.sub(
                    r"\s+",
                    " ",
                    text
                )

                return text.strip().lower()


            # ==================================================
            # FIND HEADINGS
            # ==================================================

            heading_columns = {}

            for c in range(
                1,
                ws.max_column + 1
            ):

                value = ws.cell(
                    header_row,
                    c
                ).value

                if value is not None:

                    heading_columns[
                        clean_text(value)
                    ] = c


            def get_col(name):

                return heading_columns.get(
                    clean_text(name)
                )


            # ==================================================
            # BLUE FILL
            # ==================================================

            blue_fill = PatternFill(
                fill_type="solid",
                fgColor="ADD8E6"
            )

            blue_cells = set()


            def make_blue(row, col):

                # Condition cell
                ws.cell(
                    row,
                    col
                ).fill = copy(
                    blue_fill
                )

                # Track condition cell
                blue_cells.add(
                    (row, col)
                )

                # Corresponding Symbol cell
                ws.cell(
                    row,
                    1
                ).fill = copy(
                    blue_fill
                )


            # ==================================================
            # NUMBER FUNCTION
            # ==================================================

            def get_number(value):

                if value is None:
                    return None

                if isinstance(
                    value,
                    bool
                ):
                    return None

                if isinstance(
                    value,
                    (int, float)
                ):

                    return float(value)

                text = str(value).strip()

                text = text.replace(
                    ",",
                    ""
                )

                text = text.replace(
                    "%",
                    ""
                )

                try:

                    return float(text)

                except:

                    return None


            # ==================================================
            # PARENTHESES NUMBER
            # ==================================================

            def get_parentheses_number(value):

                if value is None:
                    return None

                text = str(value)

                match = re.search(
                    r"\(\s*([-+]?\d+(?:\.\d+)?)\s*\)",
                    text
                )

                if match:

                    try:

                        return float(
                            match.group(1)
                        )

                    except:

                        return None

                return None


            # ==================================================
            # RULE 1
            # Sum I < -4
            # ==================================================

            col = get_col(
                "Sum I"
            )

            if col:

                for r in range(
                    header_row + 1,
                    ws.max_row + 1
                ):

                    value = get_number(
                        ws.cell(
                            r,
                            col
                        ).value
                    )

                    if (
                        value is not None
                        and
                        value < -4
                    ):

                        make_blue(
                            r,
                            col
                        )


            # ==================================================
            # RULE 2
            # 16> C-B / Avg.4
            # Parentheses < 0.50
            # ==================================================

            col = get_col(
                "16> C-B / Avg.4"
            )

            if col:

                for r in range(
                    header_row + 1,
                    ws.max_row + 1
                ):

                    value = ws.cell(
                        r,
                        col
                    ).value

                    parent_value = (
                        get_parentheses_number(
                            value
                        )
                    )

                    if (
                        parent_value is not None
                        and
                        parent_value < 0.50
                    ):

                        make_blue(
                            r,
                            col
                        )


            # ==================================================
            # RULE 3
            # 16< D-B / Avg.4
            #
            # Parentheses < -1
            # AND
            # Parentheses < number after 16<
            # ==================================================

            col = get_col(
                "16< D-B / Avg.4"
            )

            if col:

                for r in range(
                    header_row + 1,
                    ws.max_row + 1
                ):

                    value = ws.cell(
                        r,
                        col
                    ).value

                    if value is None:
                        continue

                    text = str(value)

                    first_match = re.search(
                        r"16<\s*([-+]?\d+(?:\.\d+)?)",
                        text,
                        re.IGNORECASE
                    )

                    parent_value = (
                        get_parentheses_number(
                            value
                        )
                    )

                    if (
                        first_match
                        and
                        parent_value is not None
                    ):

                        changing_value = float(
                            first_match.group(1)
                        )

                        if (
                            parent_value < -1
                            and
                            parent_value < changing_value
                        ):

                            make_blue(
                                r,
                                col
                            )


            # ==================================================
            # RULE 4
            # Sum O2H.10
            #
            # Values below average = BLUE
            # ==================================================

            col = get_col(
                "Sum O2H.10"
            )

            if col:

                values = []

                for r in range(
                    header_row + 1,
                    ws.max_row + 1
                ):

                    value = get_number(
                        ws.cell(
                            r,
                            col
                        ).value
                    )

                    if value is not None:

                        values.append(
                            value
                        )

                if values:

                    average_value = (
                        sum(values)
                        /
                        len(values)
                    )

                    for r in range(
                        header_row + 1,
                        ws.max_row + 1
                    ):

                        value = get_number(
                            ws.cell(
                                r,
                                col
                            ).value
                        )

                        if (
                            value is not None
                            and
                            value < average_value
                        ):

                            make_blue(
                                r,
                                col
                            )


            # ==================================================
            # RULE 5
            # Sum O2L.10
            #
            # Values below average = BLUE
            # ==================================================

            col = get_col(
                "Sum O2L.10"
            )

            if col:

                values = []

                for r in range(
                    header_row + 1,
                    ws.max_row + 1
                ):

                    value = get_number(
                        ws.cell(
                            r,
                            col
                        ).value
                    )

                    if value is not None:

                        values.append(
                            value
                        )

                if values:

                    average_value = (
                        sum(values)
                        /
                        len(values)
                    )

                    for r in range(
                        header_row + 1,
                        ws.max_row + 1
                    ):

                        value = get_number(
                            ws.cell(
                                r,
                                col
                            ).value
                        )

                        if (
                            value is not None
                            and
                            value < average_value
                        ):

                            make_blue(
                                r,
                                col
                            )


            # ==================================================
            # RULE 6
            # FIRST TWO DATED O2L COLUMNS
            #
            # Values < -1 = BLUE
            # ==================================================

            o2l_date_columns = []

            for c in range(
                1,
                ws.max_column + 1
            ):

                heading = clean_text(
                    ws.cell(
                        header_row,
                        c
                    ).value
                )

                if (
                    "o2l" in heading
                    and
                    "sum o2l.10" not in heading
                ):

                    o2l_date_columns.append(
                        c
                    )


            # FIRST TWO ONLY
            o2l_date_columns = (
                o2l_date_columns[:2]
            )


            for col in o2l_date_columns:

                for r in range(
                    header_row + 1,
                    ws.max_row + 1
                ):

                    value = get_number(
                        ws.cell(
                            r,
                            col
                        ).value
                    )

                    if (
                        value is not None
                        and
                        value < -1
                    ):

                        make_blue(
                            r,
                            col
                        )


            # ==================================================
            # RULE 7
            # 10 "-ve"
            #
            # First number before "+"
            # > 1 = BLUE
            #
            # Commas remain unchanged.
            # ==================================================

            col = get_col(
                '10 "-ve"'
            )

            if col:

                for r in range(
                    header_row + 1,
                    ws.max_row + 1
                ):

                    value = ws.cell(
                        r,
                        col
                    ).value

                    if value is None:
                        continue

                    text = str(
                        value
                    ).strip()

                    match = re.match(
                        r"\s*([-+]?\d+(?:\.\d+)?)\s*\+",
                        text
                    )

                    if match:

                        first_number = float(
                            match.group(1)
                        )

                        if first_number > 1:

                            make_blue(
                                r,
                                col
                            )


            # ==================================================
            # RULE 8
            # %Chg.1 < 0
            # ==================================================

            col = get_col(
                "%Chg.1"
            )

            if col:

                for r in range(
                    header_row + 1,
                    ws.max_row + 1
                ):

                    value = get_number(
                        ws.cell(
                            r,
                            col
                        ).value
                    )

                    if (
                        value is not None
                        and
                        value < 0
                    ):

                        make_blue(
                            r,
                            col
                        )


            # ==================================================
            # RULE 9
            # %Chg.2 < 0
            #
            # %Chg.3 and %Chg.4 NOT CHECKED
            # ==================================================

            col = get_col(
                "%Chg.2"
            )

            if col:

                for r in range(
                    header_row + 1,
                    ws.max_row + 1
                ):

                    value = get_number(
                        ws.cell(
                            r,
                            col
                        ).value
                    )

                    if (
                        value is not None
                        and
                        value < 0
                    ):

                        make_blue(
                            r,
                            col
                        )


            # ==================================================
            # COUNT BLUE CONDITION CELLS
            #
            # Column A is excluded.
            # ==================================================

            row_blue_counts = {}

            for r, c in blue_cells:

                if c == 1:
                    continue

                row_blue_counts[r] = (
                    row_blue_counts.get(
                        r,
                        0
                    ) + 1
                )


            # ==================================================
            # FIND HIGHEST BLUE COUNT
            # ==================================================

            if row_blue_counts:

                highest_blue_count = max(
                    row_blue_counts.values()
                )

            else:

                highest_blue_count = 0


            # ==================================================
            # GREEN LOGIC
            #
            # ONLY rows having HIGHEST count
            #
            # Minimum requirement = 8
            #
            # Same LIGHT GREEN for all.
            # ==================================================

            light_green_fill = PatternFill(
                fill_type="solid",
                fgColor="90EE90"
            )


            green_rows = []

            if highest_blue_count >= 8:

                for r, count in (
                    row_blue_counts.items()
                ):

                    if count == highest_blue_count:

                        # GREEN ONLY COLUMN A
                        ws.cell(
                            r,
                            1
                        ).fill = copy(
                            light_green_fill
                        )

                        green_rows.append(
                            r
                        )


            # ==================================================
            # FILTER ALL HEADINGS
            # ==================================================

            last_row = ws.max_row
            last_column = ws.max_column

            last_column_letter = (
                get_column_letter(
                    last_column
                )
            )

            ws.auto_filter.ref = (
                f"A{header_row}:"
                f"{last_column_letter}"
                f"{last_row}"
            )

            ws.freeze_panes = (
                f"A{header_row + 1}"
            )


            # ==================================================
            # SAVE TO MEMORY
            # ==================================================

            output_buffer = io.BytesIO()

            wb.save(
                output_buffer
            )

            output_buffer.seek(0)


            # ==================================================
            # REPORT
            # ==================================================

            st.success(
                "Excel file generated successfully."
            )

            col1, col2, col3 = st.columns(3)

            with col1:

                st.metric(
                    "Blue condition cells",
                    len(blue_cells)
                )

            with col2:

                st.metric(
                    "Highest blue count",
                    highest_blue_count
                )

            with col3:

                st.metric(
                    "Green rows",
                    len(green_rows)
                )


            if highest_blue_count >= 8:

                st.info(
                    f"Column A is light green for "
                    f"{len(green_rows)} row(s) having "
                    f"the highest blue count of "
                    f"{highest_blue_count}."
                )

            else:

                st.info(
                    "No green applied because the "
                    "highest blue count is below 8."
                )


            # ==================================================
            # DOWNLOAD
            # ==================================================

            st.download_button(
                label="⬇️ Download 6thsenseVardaanAutocolor.xlsx",
                data=output_buffer.getvalue(),
                file_name="6thsenseVardaanAutocolor.xlsx",
                mime=(
                    "application/vnd.openxmlformats-officedocument."
                    "spreadsheetml.sheet"
                )
            )


        except Exception as e:

            st.error(
                f"Error: {e}"
            )
```
