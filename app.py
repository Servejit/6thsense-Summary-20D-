import streamlit as st
import openpyxl
from openpyxl import load_workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from io import BytesIO


st.set_page_config(
    page_title="Excel Summary Generator",
    page_icon="📊",
    layout="wide"
)

st.title("📊 Excel Summary Generator")
st.write("Upload your Master Excel file to generate the Summary file.")


# ============================================================
# SAFE SHEET REFERENCE
# ============================================================

def ref(ws, col):
    name = ws.title.replace("'", "''")
    return f"'{name}'!{col}:{col}"


# ============================================================
# SUM I
# ============================================================

def sum_i_expression(row, sheets20):

    parts = []

    for ws in sheets20:
        parts.append(
            f'SUMIF({ref(ws,"A")},A{row},{ref(ws,"I")})'
        )

    return "=" + "+".join(parts)


# ============================================================
# C-B / AVG.4
# ============================================================

def cb_expression(ws, row):

    return (
        f'IFERROR(('
        f'XLOOKUP(A{row},{ref(ws,"A")},{ref(ws,"C")})-'
        f'XLOOKUP(A{row},{ref(ws,"A")},{ref(ws,"B")})'
        f')/'
        f'XLOOKUP(A{row},{ref(ws,"A")},{ref(ws,"B")})*100,'
        f'""'
        f')'
    )


# ============================================================
# D-B / AVG.4
# ============================================================

def db_expression(ws, row):

    return (
        f'IFERROR(('
        f'XLOOKUP(A{row},{ref(ws,"A")},{ref(ws,"D")})-'
        f'XLOOKUP(A{row},{ref(ws,"A")},{ref(ws,"B")})'
        f')/'
        f'XLOOKUP(A{row},{ref(ws,"A")},{ref(ws,"B")})*100,'
        f'""'
        f')'
    )


# ============================================================
# F-B / AVG.4
# ============================================================

def fb_expression(ws, row):

    return (
        f'IFERROR(('
        f'XLOOKUP(A{row},{ref(ws,"A")},{ref(ws,"F")})-'
        f'XLOOKUP(A{row},{ref(ws,"A")},{ref(ws,"B")})'
        f')/'
        f'XLOOKUP(A{row},{ref(ws,"A")},{ref(ws,"B")})*100,'
        f'""'
        f')'
    )


# ============================================================
# VOL.EXPAND
#
# Vol.Expand 1:
# Sheet 1 J - Average(Sheet 2,3,4 J)
#
# Vol.Expand 2:
# Sheet 2 J - Average(Sheet 3,4,5 J)
# ============================================================

def volume_expression(first_ws, average_sheets, row):

    first_volume = (
        f'SUMIF('
        f'{ref(first_ws,"A")},'
        f'A{row},'
        f'{ref(first_ws,"J")}'
        f')'
    )

    average_values = []

    for ws in average_sheets:

        average_values.append(
            f'SUMIF('
            f'{ref(ws,"A")},'
            f'A{row},'
            f'{ref(ws,"J")}'
            f')'
        )

    average_expression = (
        "(" +
        "+".join(average_values) +
        f")/{len(average_values)}"
    )

    return (
        f'=IFERROR('
        f'{first_volume}-({average_expression}),'
        f'""'
        f')'
    )


# ============================================================
# PROCESS FILE
# ============================================================

def process_excel(uploaded_file):

    input_bytes = uploaded_file.getvalue()

    wb = load_workbook(
        BytesIO(input_bytes)
    )

    # --------------------------------------------------------
    # DELETE OLD SUMMARY
    # --------------------------------------------------------

    if "summary" in wb.sheetnames:
        del wb["summary"]

    # --------------------------------------------------------
    # SHEETS
    # --------------------------------------------------------

    sheets = wb.worksheets[:]

    if len(sheets) < 20:
        raise ValueError(
            f"At least 20 sheets are required. "
            f"Found only {len(sheets)}."
        )

    first_sheet = sheets[0]

    sheets20 = sheets[:20]
    sheets10 = sheets[:10]
    sheets5 = sheets[:5]

    # --------------------------------------------------------
    # ROW RANGE
    # --------------------------------------------------------

    START_ROW = 2
    END_ROW = 211

    # --------------------------------------------------------
    # CREATE SUMMARY
    # --------------------------------------------------------

    summary = wb.create_sheet("summary")

    # --------------------------------------------------------
    # COLUMN A
    # --------------------------------------------------------

    summary["A1"] = first_sheet["A1"].value

    for r in range(START_ROW, END_ROW + 1):

        summary.cell(
            r,
            1
        ).value = first_sheet.cell(
            r,
            1
        ).value

    # --------------------------------------------------------
    # HEADINGS
    # --------------------------------------------------------

    summary["B1"] = "Sum I"
    summary["C1"] = "16> C-B / Avg.4"
    summary["D1"] = "16< D-B / Avg.4"
    summary["E1"] = "Sum O2H.10"
    summary["F1"] = "Sum O2L.10"
    summary["G1"] = "Vol.Expand 1"
    summary["H1"] = "Vol.Expand 2"

    # ========================================================
    # MAIN CALCULATIONS
    # ========================================================

    for r in range(START_ROW, END_ROW + 1):

        # ----------------------------------------------------
        # B = SUM I
        # ----------------------------------------------------

        summary.cell(
            r,
            2
        ).value = sum_i_expression(
            r,
            sheets20
        )

        # ----------------------------------------------------
        # C = 16> C-B / Avg.4
        # ----------------------------------------------------

        cb_values = [
            cb_expression(ws, r)
            for ws in sheets20
        ]

        cb_choose = (
            "CHOOSE({" +
            ",".join(
                str(i)
                for i in range(1, 21)
            ) +
            "}," +
            ",".join(cb_values) +
            ")"
        )

        summary.cell(
            r,
            3
        ).value = (
            f'=IFERROR('
            f'"16>"&TEXT('
            f'LARGE({cb_choose},17),'
            f'"0.00")'
            f'&", Avg.4 ("&'
            f'TEXT(AVERAGE('
            + ",".join(cb_values[:4])
            + '),"0.00")&")",'
            f'""'
            f')'
        )

        # ----------------------------------------------------
        # D = 16< D-B / Avg.4
        # ----------------------------------------------------

        db_values = [
            db_expression(ws, r)
            for ws in sheets20
        ]

        db_choose = (
            "CHOOSE({" +
            ",".join(
                str(i)
                for i in range(1, 21)
            ) +
            "}," +
            ",".join(db_values) +
            ")"
        )

        summary.cell(
            r,
            4
        ).value = (
            f'=IFERROR('
            f'"16<"&TEXT('
            f'SMALL({db_choose},17),'
            f'"0.00")'
            f'&", Avg.4 ("&'
            f'TEXT(AVERAGE('
            + ",".join(db_values[:4])
            + '),"0.00")&")",'
            f'""'
            f')'
        )

        # ----------------------------------------------------
        # E = SUM O2H.10
        # ----------------------------------------------------

        summary.cell(
            r,
            5
        ).value = (
            "=" +
            "+".join(
                cb_expression(ws, r)
                for ws in sheets10
            )
        )

        # ----------------------------------------------------
        # F = SUM O2L.10
        # ----------------------------------------------------

        summary.cell(
            r,
            6
        ).value = (
            "=" +
            "+".join(
                db_expression(ws, r)
                for ws in sheets10
            )
        )

        # ----------------------------------------------------
        # G = VOL.EXPAND 1
        # Sheet 1 J - Average Sheet 2,3,4 J
        # ----------------------------------------------------

        summary.cell(
            r,
            7
        ).value = volume_expression(
            sheets5[0],
            sheets5[1:4],
            r
        )

        # ----------------------------------------------------
        # H = VOL.EXPAND 2
        # Sheet 2 J - Average Sheet 3,4,5 J
        # ----------------------------------------------------

        summary.cell(
            r,
            8
        ).value = volume_expression(
            sheets5[1],
            sheets5[2:5],
            r
        )

    # ========================================================
    # O2H = I:R
    # ========================================================

    for col_num, ws in enumerate(
        sheets10,
        start=9
    ):

        summary.cell(
            1,
            col_num
        ).value = f"{ws.title} O2H"

        for r in range(
            START_ROW,
            END_ROW + 1
        ):

            summary.cell(
                r,
                col_num
            ).value = (
                "=" +
                cb_expression(
                    ws,
                    r
                )
            )

    # ========================================================
    # O2L = S:AB
    # ========================================================

    for col_num, ws in enumerate(
        sheets10,
        start=19
    ):

        summary.cell(
            1,
            col_num
        ).value = f"{ws.title} O2L"

        for r in range(
            START_ROW,
            END_ROW + 1
        ):

            summary.cell(
                r,
                col_num
            ).value = (
                "=" +
                db_expression(
                    ws,
                    r
                )
            )

    # ========================================================
    # C2O = AC:AL
    # ========================================================

    for col_num, ws in enumerate(
        sheets10,
        start=29
    ):

        summary.cell(
            1,
            col_num
        ).value = f"{ws.title} C2O"

        for r in range(
            START_ROW,
            END_ROW + 1
        ):

            summary.cell(
                r,
                col_num
            ).value = (
                "=" +
                fb_expression(
                    ws,
                    r
                )
            )

    # ========================================================
    # 10 "-ve"
    # ========================================================

    summary["AM1"] = '10 "-ve"'

    for r in range(
        START_ROW,
        END_ROW + 1
    ):

        negative_count = (
            f'IF(AC{r}<0,'
            f'IF(AD{r}<0,'
            f'IF(AE{r}<0,'
            f'IF(AF{r}<0,'
            f'IF(AG{r}<0,'
            f'IF(AH{r}<0,'
            f'IF(AI{r}<0,'
            f'IF(AJ{r}<0,'
            f'IF(AK{r}<0,'
            f'IF(AL{r}<0,10,9),'
            f'8),7),6),5),4),3),2),1),0)'
        )

        c_inside = (
            f'IFERROR('
            f'ABS(VALUE(MID('
            f'C{r},'
            f'FIND("(",C{r})+1,'
            f'FIND(")",C{r})-'
            f'FIND("(",C{r})-1'
            f'))),999999999)'
        )

        d_inside = (
            f'IFERROR('
            f'ABS(VALUE(MID('
            f'D{r},'
            f'FIND("(",D{r})+1,'
            f'FIND(")",D{r})-'
            f'FIND("(",D{r})-1'
            f'))),0)'
        )

        plus_condition = (
            f'IF('
            f'{c_inside}<'
            f'{d_inside},'
            f'"+"'
            f',"")'
        )

        positions = []

        for idx, (
            left_col,
            right_col
        ) in enumerate(
            zip(
                range(9, 19),
                range(19, 29)
            ),
            start=1
        ):

            left = get_column_letter(
                left_col
            )

            right = get_column_letter(
                right_col
            )

            positions.append(
                f'IF(AND('
                f'ISNUMBER({left}{r}),'
                f'ISNUMBER({right}{r}),'
                f'ABS({left}{r})<'
                f'ABS({right}{r})'
                f'),'
                f'"{idx}",'
                f'""'
                f')'
            )

        position_text = (
            'TEXTJOIN(",",TRUE,' +
            ",".join(positions) +
            ')'
        )

        summary.cell(
            r,
            39
        ).value = (
            f'=IFERROR('
            f'{negative_count}&'
            f'{plus_condition}&'
            f'{position_text},'
            f'""'
            f')'
        )

    # ========================================================
    # 10 "+ve"
    # ========================================================

    summary["AN1"] = '10 "+ve"'

    for r in range(
        START_ROW,
        END_ROW + 1
    ):

        positive_count = (
            f'IF(AC{r}>0,'
            f'IF(AD{r}>0,'
            f'IF(AE{r}>0,'
            f'IF(AF{r}>0,'
            f'IF(AG{r}>0,'
            f'IF(AH{r}>0,'
            f'IF(AI{r}>0,'
            f'IF(AJ{r}>0,'
            f'IF(AK{r}>0,'
            f'IF(AL{r}>0,10,9),'
            f'8),7),6),5),4),3),2),1),0)'
        )

        c_inside = (
            f'IFERROR('
            f'ABS(VALUE(MID('
            f'C{r},'
            f'FIND("(",C{r})+1,'
            f'FIND(")",C{r})-'
            f'FIND("(",C{r})-1'
            f'))),0)'
        )

        d_inside = (
            f'IFERROR('
            f'ABS(VALUE(MID('
            f'D{r},'
            f'FIND("(",D{r})+1,'
            f'FIND(")",D{r})-'
            f'FIND("(",D{r})-1'
            f'))),999999999)'
        )

        plus_condition = (
            f'IF('
            f'{c_inside}>'
            f'{d_inside},'
            f'"+"'
            f',"")'
        )

        positions = []

        for idx, (
            left_col,
            right_col
        ) in enumerate(
            zip(
                range(9, 19),
                range(19, 29)
            ),
            start=1
        ):

            left = get_column_letter(
                left_col
            )

            right = get_column_letter(
                right_col
            )

            positions.append(
                f'IF(AND('
                f'ISNUMBER({left}{r}),'
                f'ISNUMBER({right}{r}),'
                f'ABS({left}{r})>'
                f'ABS({right}{r})'
                f'),'
                f'"{idx}",'
                f'""'
                f')'
            )

        position_text = (
            'TEXTJOIN(",",TRUE,' +
            ",".join(positions) +
            ')'
        )

        summary.cell(
            r,
            40
        ).value = (
            f'=IFERROR('
            f'{positive_count}&'
            f'{plus_condition}&'
            f'{position_text},'
            f'""'
            f')'
        )

    # ========================================================
    # %Chg.1 TO %Chg.4
    # ========================================================

    summary["AO1"] = "%Chg.1"
    summary["AP1"] = "%Chg.2"
    summary["AQ1"] = "%Chg.3"
    summary["AR1"] = "%Chg.4"

    for r in range(
        START_ROW,
        END_ROW + 1
    ):

        # AO = Sheet 1 Column I
        summary.cell(
            r,
            41
        ).value = (
            f'=IFERROR('
            f'SUMIF('
            f'{ref(sheets[0],"A")},'
            f'A{r},'
            f'{ref(sheets[0],"I")}'
            f'),""'
            f')'
        )

        # AP = Sheet 2 Column I
        summary.cell(
            r,
            42
        ).value = (
            f'=IFERROR('
            f'SUMIF('
            f'{ref(sheets[1],"A")},'
            f'A{r},'
            f'{ref(sheets[1],"I")}'
            f'),""'
            f')'
        )

        # AQ = Sheet 3 Column I
        summary.cell(
            r,
            43
        ).value = (
            f'=IFERROR('
            f'SUMIF('
            f'{ref(sheets[2],"A")},'
            f'A{r},'
            f'{ref(sheets[2],"I")}'
            f'),""'
            f')'
        )

        # AR = Sheet 4 Column I
        summary.cell(
            r,
            44
        ).value = (
            f'=IFERROR('
            f'SUMIF('
            f'{ref(sheets[3],"A")},'
            f'A{r},'
            f'{ref(sheets[3],"I")}'
            f'),""'
            f')'
        )

    # ========================================================
    # FORMATTING
    # ========================================================

    summary.freeze_panes = "A2"

    summary.column_dimensions["A"].width = 22
    summary.column_dimensions["B"].width = 18
    summary.column_dimensions["C"].width = 28
    summary.column_dimensions["D"].width = 28
    summary.column_dimensions["E"].width = 18
    summary.column_dimensions["F"].width = 18
    summary.column_dimensions["G"].width = 18
    summary.column_dimensions["H"].width = 18

    for col_num in range(9, 39):
        summary.column_dimensions[
            get_column_letter(col_num)
        ].width = 16

    summary.column_dimensions["AM"].width = 30
    summary.column_dimensions["AN"].width = 30

    for col_num in range(41, 45):
        summary.column_dimensions[
            get_column_letter(col_num)
        ].width = 14

    # --------------------------------------------------------
    # NUMBER FORMATS
    # --------------------------------------------------------

    for r in range(
        START_ROW,
        END_ROW + 1
    ):

        summary.cell(
            r,
            5
        ).number_format = "0.00"

        summary.cell(
            r,
            6
        ).number_format = "0.00"

        summary.cell(
            r,
            7
        ).number_format = "+0;-0;0"

        summary.cell(
            r,
            8
        ).number_format = "+0;-0;0"

    # --------------------------------------------------------
    # ALIGNMENT
    # --------------------------------------------------------

    for row in summary.iter_rows():

        for cell in row:

            cell.alignment = Alignment(
                vertical="center"
            )

    # --------------------------------------------------------
    # AUTOFIT
    # --------------------------------------------------------

    for col_cells in summary.columns:

        max_length = 0

        column = get_column_letter(
            col_cells[0].column
        )

        for cell in col_cells:

            if cell.value is not None:

                try:
                    max_length = max(
                        max_length,
                        len(str(cell.value))
                    )
                except:
                    pass

        summary.column_dimensions[
            column
        ].width = min(
            max_length + 2,
            40
        )

    # --------------------------------------------------------
    # RESTORE IMPORTANT WIDTHS
    # --------------------------------------------------------

    summary.column_dimensions["A"].width = 22
    summary.column_dimensions["B"].width = 18
    summary.column_dimensions["C"].width = 28
    summary.column_dimensions["D"].width = 28
    summary.column_dimensions["E"].width = 18
    summary.column_dimensions["F"].width = 18
    summary.column_dimensions["G"].width = 18
    summary.column_dimensions["H"].width = 18

    for col_num in range(9, 39):
        summary.column_dimensions[
            get_column_letter(col_num)
        ].width = 16

    summary.column_dimensions["AM"].width = 30
    summary.column_dimensions["AN"].width = 30

    for col_num in range(41, 45):
        summary.column_dimensions[
            get_column_letter(col_num)
        ].width = 14

    # --------------------------------------------------------
    # FILTER
    # --------------------------------------------------------

    summary.auto_filter.ref = (
        f"A1:AR{END_ROW}"
    )

    # --------------------------------------------------------
    # FORCE RECALCULATION
    # --------------------------------------------------------

    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"

    # --------------------------------------------------------
    # SAVE TO MEMORY
    # --------------------------------------------------------

    output = BytesIO()

    wb.save(output)

    output.seek(0)

    return output


# ============================================================
# STREAMLIT UI
# ============================================================

uploaded_file = st.file_uploader(
    "Upload Master Excel File",
    type=["xlsx"]
)


if uploaded_file is not None:

    st.success(
        f"File uploaded: {uploaded_file.name}"
    )

    if st.button(
        "🚀 Generate Summary",
        type="primary"
    ):

        try:

            with st.spinner(
                "Generating summary..."
            ):

                output_file = process_excel(
                    uploaded_file
                )

            st.success(
                "✅ Summary file generated successfully."
            )

            st.download_button(
                label="⬇️ Download summary_output.xlsx",
                data=output_file,
                file_name="summary_output.xlsx",
                mime=(
                    "application/vnd.openxmlformats-officedocument."
                    "spreadsheetml.sheet"
                )
            )

        except Exception as e:

            st.error(
                f"Error: {e}"
            )
